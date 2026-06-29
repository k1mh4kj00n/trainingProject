"""우측 Parameter & Results 테이블.

NIRS-VOT 12지표 + 설정 xlsx `Coded Data` 시트의 ground truth 비교 컬럼.

표시 형식 (4 컬럼):
    Item | Ours | Ground truth | err %

ground truth 매핑:
    Baseline SmO2 (%)       ← Calf_SmO2_Base
    Min SmO2 (%)            ← Calf_SmO2_Min
    Peak SmO2 (%)           ← Calf_SmO2_Peak
    Slope1 30-150s (%/s)    ← VOT_Slope1_30_150
    Slope2 0-10s (%/s)      ← VOT_Slope2_0_10
    T50 (s)                 ← VOT_T50
    AUC 3min (%·min)        ← AUC_3min

Coded Data 의 Timepoint 매핑: session=Baseline → 'VOT_Baseline', Recovery2 → 'VOT_Post'.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
from PySide6.QtCore import Qt, Signal
from PySide6.QtGui import QColor
from PySide6.QtWidgets import (
    QFileDialog,
    QFrame,
    QGroupBox,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QMessageBox,
    QPushButton,
    QScrollArea,
    QSizePolicy,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)


_GT_COL_BY_ITEM = {
    "Baseline SmO2 (%)":     "Calf_SmO2_Base",
    "Min SmO2 (%)":          "Calf_SmO2_Min",
    "Peak SmO2 (%)":         "Calf_SmO2_Peak",
    "Slope1 30-150s (%/s)":  "VOT_Slope1_30_150",
    "Slope2 0-10s (%/s)":    "VOT_Slope2_0_10",
    "T50 (s)":               "VOT_T50",
    "AUC 3min (%·min)":      "AUC_3min",
}

# Coded Data 시트의 ground-truth 행은 'VOT_Baseline' / 'VOT_Post' 두 가지뿐.
# v0.7 동적 timepoint 시대에 매칭 정책:
#   - 이름이 정확히 'Baseline'  → VOT_Baseline 와 비교
#   - 이름이 정확히 'Recovery2' → VOT_Post 와 비교 (구버전 호환)
#   - 그 외 이름(사용자 추가 Recovery5, Post30 등) → GT 매칭 행 없음 → '—' 로 표시
# (정책 변경: vot_truth 행 자체를 사용자가 추가하려면 Coded Data 시트 확장 필요 — 별도 task)
_SESSION_TO_TIMEPOINT = {
    "Baseline":  "Baseline",
    "Recovery2": "Post",
}


class ParameterTable(QWidget):
    """Analysis Results 패널 — NIRS-VOT 지표 비교 (Ours vs Ground truth) 표기 전용.

    Auto detect Min/Peak 기능은 ``ParameterPanel`` 의 Min/Peak 그룹으로 이동 (v0.6.24).
    여기는 분석 결과 값 비교 표 + Session context (subject/anchor/detected/TTE) 표시 + Export.
    """

    export_all_requested = Signal()  # "Export all sessions" 클릭 — app.py 가 multi-session CSV 생성

    def __init__(self, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self._last_result = None
        self._last_df = None
        self._loaded_count = 0   # set_loaded_count() 로 갱신 — Export-all 활성화 조건

        # 도크 높이가 작아도 Export 버튼이 항상 보이도록:
        #   상단 (title / session label / table / context) = QScrollArea 안
        #   하단 (Export 버튼 행) = ScrollArea 밖, 도크 바닥에 고정
        outer = QVBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)
        outer.setSpacing(0)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QScrollArea.NoFrame)
        scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        scroll.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        outer.addWidget(scroll, stretch=1)

        body = QWidget()
        scroll.setWidget(body)

        root = QVBoxLayout(body)
        root.setContentsMargins(6, 6, 6, 6)

        title = QLabel("Analysis Results  ·  Ours vs Ground truth")
        title.setProperty("role", "title")
        root.addWidget(title)

        self.session_label = QLabel("— 세션을 선택하세요 —")
        self.session_label.setStyleSheet("color: #666666;")
        self.session_label.setWordWrap(True)
        root.addWidget(self.session_label)

        self.table = QTableWidget(0, 4)
        self.table.setHorizontalHeaderLabels(["Item", "Ours", "Ground truth", "err %"])
        self.table.verticalHeader().setVisible(False)
        self.table.setAlternatingRowColors(True)
        h = self.table.horizontalHeader()
        h.setSectionResizeMode(0, QHeaderView.Stretch)
        for col in (1, 2, 3):
            h.setSectionResizeMode(col, QHeaderView.ResizeToContents)
        self.table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.table.setSelectionBehavior(QTableWidget.SelectRows)
        # 12지표 정도가 한눈에 보이는 최소 높이 — 도크가 좁아지면 ScrollArea 가 처리
        self.table.setMinimumHeight(280)
        self.table.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        root.addWidget(self.table, stretch=1)

        # ── Session context (12 metrics 외 부가 정보) ──
        ctx_box = QGroupBox("Session context")
        ctx_lay = QVBoxLayout(ctx_box)
        ctx_lay.setContentsMargins(8, 12, 8, 8)
        ctx_lay.setSpacing(2)
        self.subject_line   = QLabel("Subject: —")
        self.anchor_line    = QLabel("Anchors: —")
        self.detected_line  = QLabel("Detected: —")
        self.exercise_line  = QLabel("Exercise: —")
        for lbl in (self.subject_line, self.anchor_line, self.detected_line, self.exercise_line):
            lbl.setTextFormat(Qt.RichText)
            lbl.setWordWrap(True)
            lbl.setStyleSheet("color: #24292E; padding: 1px 0;")
            ctx_lay.addWidget(lbl)
        root.addWidget(ctx_box)
        self._cfg_cached = None  # 마지막 show_result 의 session_config

        # ── Export 버튼 — ScrollArea 밖, 도크 바닥에 항상 고정 ──
        sep = QFrame()
        sep.setFrameShape(QFrame.HLine)
        sep.setStyleSheet("color: #E5E7EB;")
        sep.setFixedHeight(1)
        outer.addWidget(sep)

        btn_container = QWidget()
        btn_container.setStyleSheet("background: #FAFBFC;")
        btn_row = QHBoxLayout(btn_container)
        btn_row.setContentsMargins(6, 6, 6, 6)
        btn_row.addStretch()
        self.export_csv_btn = QPushButton("Export results…")
        self.export_csv_btn.setToolTip(
            "현재 세션 결과를 xlsx / csv / txt 중 하나로 저장.\n"
            "12 metrics + Ground truth + 메타(피험자·anchor·TTE·검출시각) 포함."
        )
        self.export_csv_btn.setEnabled(False)
        self.export_csv_btn.clicked.connect(self.export_csv)
        self.export_all_btn = QPushButton("Export all sessions…")
        self.export_all_btn.setToolTip(
            "로드된 모든 condition × session 을 wide-format (xlsx / csv / txt) 로 저장.\n"
            "각 행 = (condition, session) — Subject·metrics·GT·TTE 모두 포함."
        )
        self.export_all_btn.setEnabled(False)
        self.export_all_btn.clicked.connect(self.export_all_requested.emit)
        self.export_pdf_btn = QPushButton("Export PDF (v0.7)")
        self.export_pdf_btn.setEnabled(False)
        btn_row.addWidget(self.export_csv_btn)
        btn_row.addWidget(self.export_all_btn)
        btn_row.addWidget(self.export_pdf_btn)
        # btn_container 는 outer 에 직접 — root(scroll body) 가 아닌 도크 바닥에 고정
        outer.addWidget(btn_container)

    # ─────────────────────────────────────────────
    def show_result(self, result, df=None, *, session_config=None) -> None:
        """MetricsResult + (선택) SessionConfig 를 받아 ground truth 와 비교 렌더."""
        self._last_result = result
        self._last_df = df
        self._cfg_cached = session_config

        self.session_label.setText(
            f"{result.condition}  /  {result.session}   "
            f"({result.anchors_abs['s'].strftime('%H:%M:%S')} → "
            f"{result.anchors_abs['e'].strftime('%H:%M:%S')})"
        )

        gt_row = _resolve_gt_row(session_config, result.session)

        tbl = result.as_table()
        rows = list(tbl.itertuples(index=False))
        self.table.setRowCount(len(rows))
        for r, (item, val) in enumerate(rows):
            gt_val = _lookup_gt(gt_row, str(item))
            self.table.setItem(r, 0, _name_item(str(item)))
            self.table.setItem(r, 1, _val_item(val))

            if gt_val is None:
                self.table.setItem(r, 2, _val_item(float("nan")))
                self.table.setItem(r, 3, _err_item(float("nan")))
                continue
            try:
                gt_f = float(gt_val)
            except (TypeError, ValueError):
                # Excel 입력 오류 등 비숫자 값
                gt_item = QTableWidgetItem(str(gt_val))
                gt_item.setForeground(QColor("#888888"))
                gt_item.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                self.table.setItem(r, 2, gt_item)
                self.table.setItem(r, 3, _err_item(float("nan")))
                continue

            err = (
                abs(float(val) - gt_f) / max(abs(gt_f), 1e-9) * 100
                if gt_f != 0
                else (0.0 if val == 0 else float("inf"))
            )
            self.table.setItem(r, 2, _val_item(gt_f))
            self.table.setItem(r, 3, _err_item(err))

        # Session context — subject / anchor / detected / TTE
        self._render_context(result, session_config)

        self.export_csv_btn.setEnabled(True)

    def _render_context(self, result, cfg) -> None:
        """Session context 4 라벨 갱신 — show_result 직후 호출."""
        # Subject
        info = (cfg.subject_info if cfg is not None else {}) or {}
        subj_name = info.get("Name") or info.get("ID(Protocol)") or "—"
        parts = [f"<b>{subj_name}</b>"]
        if (age := info.get("Age")) is not None and age != 0:
            gender = info.get("Gender") or ""
            short_g = {"Male": "M", "Female": "F"}.get(str(gender), str(gender)[:1] if gender else "")
            parts.append(f"{age}{short_g}")
        wt = info.get("Weight(kg)"); ht = info.get("Height(cm)")
        if wt or ht:
            parts.append(f"{ht or '—'}cm / {wt or '—'}kg")
        wmax = info.get("Wattmax"); wkg = info.get("W/kg")
        if wmax:
            extra = f"Wattmax {wmax}W"
            if wkg:
                extra += f" ({wkg} W/kg)"
            parts.append(extra)
        if (vo2 := info.get("VO2max")):
            parts.append(f"VO2max {vo2}")
        self.subject_line.setText("Subject: " + " · ".join(parts))

        # Anchors (s/i/d/e clock)
        a = result.anchors_abs
        self.anchor_line.setText(
            f"Anchors: <b>{a['s'].strftime('%H:%M:%S')}</b> → "
            f"{a['i'].strftime('%H:%M:%S')} (Occ) → "
            f"{a['d'].strftime('%H:%M:%S')} (Def) → "
            f"<b>{a['e'].strftime('%H:%M:%S')}</b>"
        )

        # Detected min/peak
        mt = getattr(result, "min_time", None)
        pt = getattr(result, "peak_time", None)
        mv = getattr(result, "min_smo2", float("nan"))
        pv = getattr(result, "peak_smo2", float("nan"))
        det_min = (
            f"<span style='color:#DC2626;'>Min @ {mt.strftime('%H:%M:%S')}</span> = "
            f"<b>{mv:.2f}%</b>"
            if mt is not None and np.isfinite(mv) else "Min @ — = —"
        )
        det_peak = (
            f"<span style='color:#16A34A;'>Peak @ {pt.strftime('%H:%M:%S')}</span> = "
            f"<b>{pv:.2f}%</b>"
            if pt is not None and np.isfinite(pv) else "Peak @ — = —"
        )
        mag = getattr(result, "magnitude", float("nan"))
        mag_s = f"  ·  Magnitude <b>{mag:.2f}%</b>" if np.isfinite(mag) else ""
        self.detected_line.setText(f"Detected: {det_min}  ·  {det_peak}{mag_s}")

        # Exercise (TTE)
        ex = (cfg.timepoints.get("exercise", {}) or {}) if cfg else {}
        tte1_s = _time_to_sec(ex.get("TTE1"))
        tte2_s = _time_to_sec(ex.get("TTE2"))
        maint = ex.get("TTE_maintenance_rate")
        try:
            maint_f = float(maint) if maint is not None else 0.0
        except (TypeError, ValueError):
            maint_f = 0.0
        # 품질 컬러
        if maint_f <= 0:
            mlabel = "—"
        elif maint_f < 70.0:
            mlabel = (
                f"<span style='color:#DC2626; font-weight:bold;'>{maint_f:.1f}%</span> "
                "<span style='color:#6A737D;'>[&lt; 70% 경고]</span>"
            )
        else:
            mlabel = (
                f"<span style='color:#16A34A; font-weight:bold;'>{maint_f:.1f}%</span> "
                "<span style='color:#6A737D;'>[≥ 70%]</span>"
            )
        tte1_str = f"{tte1_s}s" if tte1_s > 0 else "—"
        tte2_str = f"{tte2_s}s" if tte2_s > 0 else "—"
        self.exercise_line.setText(
            f"Exercise: TTE1 <b>{tte1_str}</b>  ·  TTE2 <b>{tte2_str}</b>  ·  Maintenance {mlabel}"
        )

    def set_loaded_count(self, n: int) -> None:
        """app.py 가 NirsVotTab.loaded 의 condition 수를 알려옴 — Export-all 버튼 활성화."""
        self._loaded_count = max(0, int(n))
        self.export_all_btn.setEnabled(self._loaded_count > 0)

    def clear(self) -> None:
        self.table.setRowCount(0)
        self.session_label.setText("— 세션을 선택하세요 —")
        self._last_result = None
        self._cfg_cached = None
        self.export_csv_btn.setEnabled(False)
        self.subject_line.setText("Subject: —")
        self.anchor_line.setText("Anchors: —")
        self.detected_line.setText("Detected: —")
        self.exercise_line.setText("Exercise: —")

    def export_csv(self) -> None:
        """현재 세션의 metrics + GT + 메타데이터를 xlsx/csv/txt 중 하나로 저장.

        포맷:
          - 상단: 메타데이터 (subject, anchors, detected, TTE)
          - 본문: Item / Ours / Ground truth / err % 4 컬럼
                   12 metrics + 추가 행 (Min time, Peak time, TTE1, TTE2, Maintenance)
        확장자(.xlsx / .csv / .txt) 로 포맷 자동 결정.
        """
        if self._last_result is None:
            return
        r = self._last_result
        cfg = self._cfg_cached
        default = f"{r.condition}_{r.session}_results.xlsx"
        path, selected = QFileDialog.getSaveFileName(
            self, "Export results", default,
            "Excel (*.xlsx);;CSV (*.csv);;Text (*.txt)",
        )
        if not path:
            return
        p = Path(path)
        # 확장자 없을 때 필터 기반 보강
        if p.suffix == "":
            if "xlsx" in selected:
                p = p.with_suffix(".xlsx")
            elif "txt" in selected:
                p = p.with_suffix(".txt")
            else:
                p = p.with_suffix(".csv")

        ext = p.suffix.lower()
        try:
            if ext in (".xlsx", ".xls"):
                self._write_single_session_xlsx(p, r, cfg)
            elif ext == ".txt":
                self._write_single_session_txt(p, r, cfg)
            else:
                self._write_single_session_csv(p, r, cfg)
        except Exception as exc:  # noqa: BLE001
            QMessageBox.critical(self, "Export 실패", str(exc))
            return
        QMessageBox.information(self, "Export", f"저장 완료:\n{p}")

    # ─────────────────────────────────────────────
    #  단일 세션 — 메타/데이터 빌더 (포맷 무관)
    # ─────────────────────────────────────────────
    @staticmethod
    def _build_meta_and_data(result, cfg) -> tuple[list[str], list[tuple[str, str, str, str]]]:
        """``(meta_lines, data_rows)`` 반환 — 모든 포맷 writer 가 공통으로 사용.

        ``meta_lines`` — 7~8 줄 메타 (각 줄 ``# `` prefix 포함).
        ``data_rows``  — ``(item, ours, gt, err)`` 4-튜플 리스트. 모두 문자열.
        """
        info = (cfg.subject_info if cfg is not None else {}) or {}
        ex   = (cfg.timepoints.get("exercise", {}) if cfg is not None else {}) or {}
        gt   = _resolve_gt_row(cfg, result.session)

        a = result.anchors_abs
        meta_lines = [
            f"# Subject: {info.get('Name', '—')} · ID {info.get('ID(Protocol)', '—')} · "
            f"Age {info.get('Age', '—')} · Gender {info.get('Gender', '—')}",
            f"# Body: {info.get('Height(cm)', '—')}cm · {info.get('Weight(kg)', '—')}kg",
            f"# Performance: Wattmax {info.get('Wattmax', '—')}W · "
            f"VO2max {info.get('VO2max', '—')} · W/kg {info.get('W/kg', '—')} · "
            f"FTP {info.get('FTP', '—')}",
            f"# Condition: {result.condition}  Session: {result.session}",
            f"# Anchors (clock): start={a['s'].strftime('%H:%M:%S')} "
            f"inflate={a['i'].strftime('%H:%M:%S')} "
            f"deflate={a['d'].strftime('%H:%M:%S')} "
            f"end={a['e'].strftime('%H:%M:%S')}",
        ]
        mt = getattr(result, "min_time", None)
        pt = getattr(result, "peak_time", None)
        mv = getattr(result, "min_smo2", float("nan"))
        pv = getattr(result, "peak_smo2", float("nan"))
        if np.isfinite(mv) and np.isfinite(pv):
            meta_lines.append(
                f"# Detected: Min @ {mt.strftime('%H:%M:%S') if mt is not None else '—'} = "
                f"{mv:.2f}%  ·  Peak @ {pt.strftime('%H:%M:%S') if pt is not None else '—'} = "
                f"{pv:.2f}%"
            )
        else:
            meta_lines.append("# Detected: —")

        tte1_s = _time_to_sec(ex.get("TTE1"))
        tte2_s = _time_to_sec(ex.get("TTE2"))
        maint = ex.get("TTE_maintenance_rate")
        try:
            mf = float(maint) if maint is not None else 0.0
        except (TypeError, ValueError):
            mf = 0.0
        flag = "✓ ≥70%" if mf >= 70.0 else ("⚠ <70%" if mf > 0 else "—")
        meta_lines.append(
            f"# Exercise: TTE1={tte1_s}s · TTE2={tte2_s}s · Maintenance={mf:.2f}% [{flag}]"
        )

        # 데이터 행 (모두 문자열)
        data_rows: list[tuple[str, str, str, str]] = []
        tbl = result.as_table()
        for item, val in tbl.itertuples(index=False):
            gv = _lookup_gt(gt, str(item))
            ours = _fmt_num(val)
            if gv is None:
                data_rows.append((str(item), ours, "", ""))
                continue
            try:
                gf = float(gv)
                err = (
                    abs(float(val) - gf) / max(abs(gf), 1e-9) * 100
                    if gf != 0 else (0.0 if val == 0 else float("inf"))
                )
                data_rows.append((str(item), ours, _fmt_num(gf), _fmt_num(err)))
            except (TypeError, ValueError):
                data_rows.append((str(item), ours, str(gv), ""))

        # 추가 행 — time-based / TTE (no GT)
        extra = [
            ("Min time (clock)",
             mt.strftime("%H:%M:%S") if mt is not None else "—"),
            ("Peak time (clock)",
             pt.strftime("%H:%M:%S") if pt is not None else "—"),
            ("TTE1 (sec)",        str(tte1_s) if tte1_s > 0 else ""),
            ("TTE2 (sec)",        str(tte2_s) if tte2_s > 0 else ""),
            ("Maintenance (%)",   _fmt_num(mf) if mf > 0 else ""),
        ]
        for item, ours in extra:
            data_rows.append((item, str(ours), "", ""))

        return meta_lines, data_rows

    @staticmethod
    def _write_single_session_csv(path: Path, result, cfg) -> None:
        """단일 세션 CSV — 메타 헤더 (# 코멘트) + 4 컬럼 데이터."""
        import csv
        meta, rows = ParameterTable._build_meta_and_data(result, cfg)
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            for line in meta:
                f.write(line + "\n")
            f.write("#\n")
            w = csv.writer(f)
            w.writerow(["Item", "Ours", "Ground truth", "err %"])
            for r in rows:
                w.writerow(r)

    @staticmethod
    def _write_single_session_xlsx(path: Path, result, cfg) -> None:
        """단일 세션 xlsx — 메타 텍스트 (A 컬럼) + 표 + 색상.

        - 메타 7~8 줄: 컬럼 A 에 한 줄씩, bold
        - 빈 행 1개
        - 헤더 행: A=Item, B=Ours, C=Ground truth, D=err %  (회색 배경, bold)
        - 데이터 행: err % 가 작으면 초록·중간 노랑·큰 빨강
        - 컬럼 폭 자동 조정
        """
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment

        meta, rows = ParameterTable._build_meta_and_data(result, cfg)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Results"

        bold = Font(bold=True)
        head_fill = PatternFill("solid", fgColor="E5E7EB")
        right = Alignment(horizontal="right")

        for line in meta:
            ws.append([line])
            ws.cell(row=ws.max_row, column=1).font = bold
        ws.append([])  # 빈 행
        # 헤더
        ws.append(["Item", "Ours", "Ground truth", "err %"])
        for c in range(1, 5):
            cell = ws.cell(row=ws.max_row, column=c)
            cell.font = bold
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center")
        # 데이터
        for item, ours, gt, err in rows:
            ws.append([item, _xnum(ours), _xnum(gt), _xnum(err)])
            r = ws.max_row
            ws.cell(row=r, column=2).alignment = right
            ws.cell(row=r, column=3).alignment = right
            ws.cell(row=r, column=4).alignment = right
            # err % 색상
            try:
                e = float(err) if err else None
                if e is not None and np.isfinite(e):
                    if e < 3:
                        color = "2E7D32"
                    elif e < 15:
                        color = "FFA000"
                    else:
                        color = "C62828"
                    ws.cell(row=r, column=4).font = Font(bold=True, color=color)
            except (TypeError, ValueError):
                pass
        # 컬럼 폭
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 12
        ws.column_dimensions["C"].width = 14
        ws.column_dimensions["D"].width = 10
        wb.save(path)

    @staticmethod
    def _write_single_session_txt(path: Path, result, cfg) -> None:
        """단일 세션 txt — 사람 읽기 좋은 fixed-width 표.

        - 메타 줄 그대로
        - 빈 줄
        - 헤더 행 + 구분선
        - 컬럼 폭 자동: Item 28, Ours 12, GT 14, err 10
        """
        meta, rows = ParameterTable._build_meta_and_data(result, cfg)
        widths = (28, 12, 14, 10)
        header = ("Item", "Ours", "Ground truth", "err %")

        def _row_str(values, sep="  "):
            cells = []
            for i, v in enumerate(values):
                w = widths[i]
                if i == 0:
                    cells.append(str(v).ljust(w))
                else:
                    cells.append(str(v).rjust(w))
            return sep.join(cells)

        lines = list(meta) + ["#", ""]
        lines.append(_row_str(header))
        lines.append("─" * (sum(widths) + 2 * (len(widths) - 1)))
        for r in rows:
            lines.append(_row_str(r))
        path.write_text("\n".join(lines) + "\n", encoding="utf-8")

    # ─────────────────────────────────────────────
    #  Multi-session — 컬럼 순서 / 디스패처
    # ─────────────────────────────────────────────
    _MULTI_PRIORITY: tuple[str, ...] = (
        "condition", "session",
        "subject_name", "subject_id", "age", "gender",
        "height_cm", "weight_kg", "wattmax", "vo2max", "w_per_kg", "ftp",
        "anchor_start", "anchor_inflate", "anchor_deflate", "anchor_end",
        "baseline_smo2", "slope1_0_60", "slope1_30_150", "oxy_deficit",
        "min_smo2", "peak_smo2", "magnitude",
        "slope2_0_10", "slope2_0_30", "t50_sec", "t95_sec", "auc_3min",
        "min_time", "peak_time",
        "tte1_sec", "tte2_sec", "maintenance_rate", "maintenance_flag",
        "gt_baseline_smo2", "gt_min_smo2", "gt_peak_smo2",
        "gt_slope1_30_150", "gt_slope2_0_10", "gt_t50", "gt_auc_3min",
    )

    @staticmethod
    def _multi_columns(rows: list[dict]) -> list[str]:
        all_keys: set[str] = set()
        for d in rows:
            all_keys.update(d.keys())
        ordered = [k for k in ParameterTable._MULTI_PRIORITY if k in all_keys]
        ordered += sorted(k for k in all_keys if k not in ParameterTable._MULTI_PRIORITY)
        return ordered

    @staticmethod
    def write_multi_session(path: Path, rows: list[dict]) -> None:
        """확장자(.xlsx/.csv/.txt)로 포맷 자동 결정해 multi-session 파일 작성."""
        ext = path.suffix.lower()
        if ext in (".xlsx", ".xls"):
            ParameterTable.write_multi_session_xlsx(path, rows)
        elif ext == ".txt":
            ParameterTable.write_multi_session_txt(path, rows)
        else:
            ParameterTable.write_multi_session_csv(path, rows)

    @staticmethod
    def write_multi_session_csv(path: Path, rows: list[dict]) -> None:
        """multi-session wide CSV — dict 리스트를 그대로 직렬화."""
        import csv
        if not rows:
            return
        ordered = ParameterTable._multi_columns(rows)
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=ordered, extrasaction="ignore")
            w.writeheader()
            for d in rows:
                w.writerow(d)

    @staticmethod
    def write_multi_session_xlsx(path: Path, rows: list[dict]) -> None:
        """multi-session wide xlsx — 헤더 bold, maintenance_flag 색상 (OK 초록 / WARN 빨강)."""
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        if not rows:
            return
        ordered = ParameterTable._multi_columns(rows)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sessions"
        ws.append(ordered)
        for c, _ in enumerate(ordered, start=1):
            cell = ws.cell(row=1, column=c)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E5E7EB")
            cell.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "C2"  # condition / session 두 컬럼 + 첫 행 고정

        flag_idx = ordered.index("maintenance_flag") + 1 if "maintenance_flag" in ordered else None
        for d in rows:
            ws.append([d.get(k, "") for k in ordered])
            r = ws.max_row
            if flag_idx is not None:
                v = str(d.get("maintenance_flag", "") or "")
                cell = ws.cell(row=r, column=flag_idx)
                if v.startswith("WARN"):
                    cell.font = Font(bold=True, color="C62828")
                elif v == "OK":
                    cell.font = Font(bold=True, color="2E7D32")
        # 컬럼 폭 — 적당히
        for c, k in enumerate(ordered, start=1):
            ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = max(
                10, min(22, len(k) + 4)
            )
        wb.save(path)

    @staticmethod
    def write_multi_session_txt(path: Path, rows: list[dict]) -> None:
        """multi-session wide TSV — 탭 구분 (Excel·pandas 양쪽 호환)."""
        if not rows:
            return
        ordered = ParameterTable._multi_columns(rows)
        with open(path, "w", encoding="utf-8") as f:
            f.write("# Physical Analysis — Multi-session results (TSV / tab-separated)\n")
            f.write(f"# {len(rows)} rows × {len(ordered)} columns\n#\n")
            f.write("\t".join(ordered) + "\n")
            for d in rows:
                f.write("\t".join(str(d.get(k, "")) for k in ordered) + "\n")


# ─────────────────────────────────────────────────────────────
#  내부 helpers
# ─────────────────────────────────────────────────────────────
def _resolve_gt_row(session_config, session: str) -> dict:
    """SessionConfig 에서 vot_truth 의 해당 timepoint dict 를 가져온다."""
    if session_config is None:
        return {}
    timepoint = _SESSION_TO_TIMEPOINT.get(session, session)
    return session_config.vot_truth.get(timepoint, {}) or {}


def _lookup_gt(gt_row: dict, item_name: str):
    col = _GT_COL_BY_ITEM.get(item_name)
    if col is None:
        return None
    return gt_row.get(col)


def _name_item(text: str) -> QTableWidgetItem:
    it = QTableWidgetItem(text)
    it.setFlags(it.flags() ^ Qt.ItemIsEditable)
    return it


def _val_item(v: float) -> QTableWidgetItem:
    if not np.isfinite(v):
        it = QTableWidgetItem("—")
    else:
        it = QTableWidgetItem(f"{v:,.4f}")
    it.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
    it.setFlags(it.flags() ^ Qt.ItemIsEditable)
    return it


def _time_to_sec(t) -> int:
    """dt.time / int / None → 초 단위 정수."""
    import datetime as _dt
    if t is None:
        return 0
    if isinstance(t, _dt.time):
        return t.hour * 3600 + t.minute * 60 + t.second
    if isinstance(t, _dt.timedelta):
        return int(t.total_seconds())
    if isinstance(t, (int, float)):
        return max(0, int(t))
    return 0


def _fmt_num(v) -> str:
    """숫자를 CSV 출력용 문자열로 — 빈/NaN 은 빈 문자열."""
    try:
        f = float(v)
    except (TypeError, ValueError):
        return str(v) if v is not None else ""
    if not np.isfinite(f):
        return ""
    if f == int(f):
        return f"{int(f)}"
    return f"{f:.4f}"


def _xnum(s):
    """문자열을 가능하면 float 로 변환해 xlsx 셀에 넣기 (Excel 이 숫자로 인식)."""
    if s == "" or s is None:
        return ""
    try:
        f = float(s)
        return f
    except (TypeError, ValueError):
        return str(s)


def _err_item(err: float) -> QTableWidgetItem:
    if not np.isfinite(err):
        it = QTableWidgetItem("—")
        it.setForeground(QColor("#888888"))
    else:
        if err < 3:
            color = QColor("#2E7D32"); sym = "✓"
        elif err < 15:
            color = QColor("#FFA000"); sym = "≈"
        else:
            color = QColor("#C62828"); sym = "✗"
        it = QTableWidgetItem(f"{sym} {err:.2f}")
        it.setForeground(color)
    it.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
    it.setFlags(it.flags() ^ Qt.ItemIsEditable)
    return it
